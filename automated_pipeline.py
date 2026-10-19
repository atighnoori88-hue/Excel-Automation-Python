import os
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill

def process_and_calculate_data(input_data):
    """Pandas Section: Create DataFrame and calculate financial metrics."""
    df = pd.DataFrame(input_data)
    
    # Calculate Total Income for each product row
    df["Total_Income"] = df["Price"] * df["Quantity"]
    
    # Calculate summary metrics
    grand_total = df["Total_Income"].sum()
    average_product_price = df["Price"].mean()
    
    # Append total and average rows using .loc
    df.loc[len(df)] = {"Product": "GRAND TOTAL", "Price": "-", "Quantity": "-", "Customer": "-", "Total_Income": grand_total}
    df.loc[len(df)] = {"Product": "AVERAGE PRICE", "Price": average_product_price, "Quantity": "-", "Customer": "-", "Total_Income": "-"}
    
    return df

def apply_visual_styles(file_path):
    """OpenPyXL Section: Load generated workbook and apply visual formatting."""
    if not os.path.exists(file_path):
        print(f"❌ Error: The file {file_path} was not found.")
        return
        
    workbook = openpyxl.load_workbook(file_path)
    sheet = workbook.active
    max_row = sheet.max_row
    
    # Define styles (Bold font and soft green background fill)
    bold_font = Font(name="Arial", size=11, bold=True)
    highlight_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
    
    # Style the GRAND TOTAL row (second to last row)
    for col in range(1, sheet.max_column + 1):
        cell = sheet.cell(row=max_row - 1, column=col)
        cell.font = bold_font
        cell.fill = highlight_fill
        
    # Style the AVERAGE PRICE row (last row)
    for col in range(1, sheet.max_column + 1):
        cell = sheet.cell(row=max_row, column=col)
        cell.font = bold_font
        
    # Save the finalized stylized workbook
    final_output = "stylized_sales_report.xlsx"
    workbook.save(final_output)
    print(f"🎯 Visual styles applied successfully to: {final_output}")

# --- Data Pipeline Execution ---
if __name__ == "__main__":
    try:
        # Input raw sales data
        raw_sales_data = [
            {"Product": "Laptop", "Price": 800, "Quantity": 3, "Customer": "John"},
            {"Product": "Mouse", "Price": 25, "Quantity": 10, "Customer": "Alice"},
            {"Product": "Keyboard", "Price": 45, "Quantity": 5, "Customer": "Bob"},
            {"Product": "Monitor", "Price": 150, "Quantity": 2, "Customer": "David"}
        ]
        
        print("🚀 Starting data automation pipeline...")
        
        # Step 1: Process with Pandas
        processed_df = process_and_calculate_data(raw_sales_data)
        temp_file = "sales_report.xlsx"
        processed_df.to_excel(temp_file, index=False)
        print("✅ Step 1: Data processed and temporary excel file generated.")
        
        # Step 2: Format with OpenPyXL
        apply_visual_styles(temp_file)
        
        print("✨ Project Success: The entire data pipeline completed successfully!")
        
    except Exception as e:
        print(f"❌ An error occurred during pipeline execution: {e}")