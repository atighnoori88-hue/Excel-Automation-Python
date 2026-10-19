import openpyxl
from openpyxl.styles import Font, PatternFill

workbook = openpyxl.load_workbook("sales_report.xlsx")
sheet = workbook.active

bold_font = Font(name="Arial", size=11, bold=True)
highlight_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid") 

max_row = sheet.max_row

for col in range(1, sheet.max_column + 1):
    cell = sheet.cell(row=max_row - 1, column=col)
    cell.font = bold_font
    cell.fill = highlight_fill

for col in range(1, sheet.max_column + 1):
    cell = sheet.cell(row=max_row, column=col)
    cell.font = bold_font

workbook.save("stylized_sales_report.xlsx")
print("✨ Visual styles and clean formatting applied successfully to the Excel sheet!")
